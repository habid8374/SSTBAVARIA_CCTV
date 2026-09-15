"""Carga masiva de Trabajador desde un Excel — pensado para el "Excel
masivo" que ya maneja el cliente: cada fila es un trabajador y una
columna "Contratista" (desplegable) indica a qué EmpresaContratista
pertenece esa fila, para poder mezclar personal de varias empresas
contratistas en un mismo archivo.

A diferencia del importador de Declaración de Método (que solo
precarga un formulario), este SÍ crea registros en la base — una fila
por Trabajador válido — reutilizando TrabajadorSerializer y sus mismas
reglas que la creación manual desde el formulario (autorización de
datos obligatoria, documento único por contratista, etc.): una fila
que no pasaría el formulario normal tampoco se crea acá, queda
reportada como error de esa fila sin tumbar el resto del archivo.

La plantilla trae, además de los datos básicos, una columna de fecha
por cada curso del catálogo Safety Academy y por cada certificación
especial (espacios confinados, conducción, manlift, grúa, soldador,
rescatista, licencia SST, ...) — para poder verificar de una vez la
vigencia de cada una al importar, igual que si se marcaran uno por uno
en el formulario manual. Ambos catálogos se leen de la base en el
momento de generar/leer el Excel (no de una lista fija en código), así
que si un Administrador agrega o desactiva un curso/certificación
desde Sistema → Reglas de contratistas, la plantilla lo refleja sin
tocar código.

La radicación de seguridad social nunca se crea desde acá — cada
trabajador importado queda igual que si se hubiera registrado a mano:
sin radicaciones, pendiente de que el contratista o SST la cargue
después (ver RadicacionSeguridadSocial, un modelo aparte)."""

import datetime
import unicodedata

import openpyxl
from openpyxl.styles import Protection
from openpyxl.utils import get_column_letter
from openpyxl.workbook.protection import WorkbookProtection
from openpyxl.worksheet.datavalidation import DataValidation

from .models import CertificacionEspecial, CursoSafetyAcademy, EmpresaContratista, Trabajador


class ErrorImportacionExcel(Exception):
    """Mensaje pensado para mostrarse tal cual al usuario."""


HOJA_TRABAJADORES = "Trabajadores"
HOJA_LISTAS = "Listas (no borrar)"

# Protege la hoja/libro contra ediciones accidentales de las fórmulas de
# validación y contra desocultar "Listas (no borrar)" — no es una medida de
# seguridad real (la protección de Excel se quita sin la contraseña con
# herramientas de terceros), solo evita que alguien la rompa sin querer.
CONTRASENA_PLANTILLA = "SSTBavaria2026"

ENCABEZADOS_BASE = [
    "Contratista",
    "Nombres",
    "Apellidos",
    "Documento",
    "Autorización datos (SI/NO)",
    "EPS",
    "ARL",
    "AFP",
    "Tipo de vinculación (Fijo/Temporal)",
    "Fecha inicio contrato (AAAA-MM-DD)",
    "Vencimiento examen médico alturas (AAAA-MM-DD)",
    "Vencimiento certificación alturas (AAAA-MM-DD)",
]

# Filas con el desplegable ya aplicado en la plantilla descargable, para
# poder pegar/escribir varias de una vez sin tener que repetir la validación.
FILAS_PLANTILLA = 300

_ETIQUETA_A_TIPO_VINCULACION = {
    "fijo": Trabajador.TipoVinculacion.FIJO,
    "temporal": Trabajador.TipoVinculacion.TEMPORAL,
}
_VALORES_SI = {"si", "s", "x", "yes", "true", "1"}


def _sin_acentos(texto):
    texto = unicodedata.normalize("NFKD", str(texto))
    return "".join(c for c in texto if not unicodedata.combining(c)).strip().lower()


def _texto_celda(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _valor_fecha(celda):
    """Acepta tanto una fecha real de Excel (la plantilla no fuerza
    formato de celda) como texto AAAA-MM-DD o DD/MM/AAAA. Una fecha
    ilegible se ignora en vez de tumbar la fila — el campo es opcional."""
    if celda in (None, ""):
        return None
    if isinstance(celda, datetime.datetime):
        return celda.date()
    if isinstance(celda, datetime.date):
        return celda
    texto = str(celda).strip()
    if not texto:
        return None
    for formato in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    return None


def _catalogo_cursos():
    return list(CursoSafetyAcademy.objects.filter(activo=True).order_by("orden", "etiqueta"))


def _catalogo_certificaciones():
    return list(CertificacionEspecial.objects.filter(activo=True).order_by("orden", "etiqueta"))


def _encabezados_y_catalogos():
    """(encabezados, claves_cursos, claves_certificaciones) — se calcula en
    cada llamada (no una vez al importar el módulo) porque ambos catálogos
    son editables desde el dashboard."""
    cursos = _catalogo_cursos()
    certificaciones = _catalogo_certificaciones()
    encabezados_cursos = [f"Curso: {c.etiqueta} — vencimiento (AAAA-MM-DD)" for c in cursos]
    encabezados_certificaciones = [
        f"Certificación: {c.etiqueta} — vencimiento (AAAA-MM-DD)" for c in certificaciones
    ]
    encabezados = ENCABEZADOS_BASE + encabezados_cursos + encabezados_certificaciones
    claves_cursos = [c.clave for c in cursos]
    claves_certificaciones = [c.clave for c in certificaciones]
    return encabezados, claves_cursos, claves_certificaciones


def generar_plantilla_trabajadores_excel():
    """Libro .xlsx en blanco con los encabezados y los desplegables
    (Contratista, autorización de datos, tipo de vinculación — solo
    empresas activas) ya aplicados a las primeras FILAS_PLANTILLA filas —
    lista para llenar y volver a subir."""
    contratistas = list(
        EmpresaContratista.objects.filter(activa=True).order_by("nombre").values_list("nombre", flat=True)
    )
    encabezados, claves_cursos, claves_certificaciones = _encabezados_y_catalogos()

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = HOJA_TRABAJADORES
    hoja.append(encabezados)
    for celda in hoja[1]:
        celda.font = celda.font.copy(bold=True)

    hoja_listas = libro.create_sheet(HOJA_LISTAS)
    for fila, nombre in enumerate(contratistas, start=1):
        hoja_listas.cell(row=fila, column=1, value=nombre)
    hoja_listas.sheet_state = "hidden"

    ultima_fila_listas = max(len(contratistas), 1)
    dv_contratista = DataValidation(
        type="list",
        formula1=f"'{HOJA_LISTAS}'!$A$1:$A${ultima_fila_listas}",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="Contratista no válido",
        error="Elige una empresa contratista de la lista desplegable.",
    )
    hoja.add_data_validation(dv_contratista)
    dv_contratista.add(f"A2:A{FILAS_PLANTILLA + 1}")

    dv_autorizacion = DataValidation(type="list", formula1='"SI,NO"', allow_blank=True)
    hoja.add_data_validation(dv_autorizacion)
    dv_autorizacion.add(f"E2:E{FILAS_PLANTILLA + 1}")

    dv_vinculacion = DataValidation(type="list", formula1='"Fijo,Temporal"', allow_blank=True)
    hoja.add_data_validation(dv_vinculacion)
    dv_vinculacion.add(f"I2:I{FILAS_PLANTILLA + 1}")

    anchos_base = [30, 20, 20, 16, 24, 14, 14, 14, 26, 24, 30, 28]
    anchos = anchos_base + [30] * (len(claves_cursos) + len(claves_certificaciones))
    for indice, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho

    # Deja libres para escribir solo las celdas de datos (encabezado incluido
    # como referencia, pero bloqueado); todo lo demás de la hoja —fórmulas de
    # validación, formato— queda protegido una vez se activa hoja.protection.
    celda_desbloqueada = Protection(locked=False)
    ultima_columna = len(encabezados)
    for fila in hoja.iter_rows(min_row=2, max_row=FILAS_PLANTILLA + 1, min_col=1, max_col=ultima_columna):
        for celda in fila:
            celda.protection = celda_desbloqueada

    hoja.protection.sheet = True
    hoja.protection.password = CONTRASENA_PLANTILLA
    hoja.protection.formatColumns = False
    hoja.protection.formatRows = False

    hoja_listas.protection.sheet = True
    hoja_listas.protection.password = CONTRASENA_PLANTILLA

    libro.security = WorkbookProtection(lockStructure=True)
    libro.security.set_workbook_password(CONTRASENA_PLANTILLA)

    return libro


def procesar_trabajadores_excel(archivo):
    """Lee el Excel y devuelve (filas, errores_lectura):

    - `filas`: lista de dicts, uno por fila con datos, con las mismas
      claves que TrabajadorSerializer espera, más "_fila_excel" (el
      número de fila real del Excel, para señalarlo en los mensajes de
      error del caller si la fila no pasa la validación del serializer).
    - `errores_lectura`: filas que no se pudieron ni armar (ej. el
      nombre de la columna Contratista no calza con ninguna empresa) —
      errores de estructura, no de validación de negocio."""
    try:
        libro = openpyxl.load_workbook(archivo, data_only=True)
    except Exception as exc:
        raise ErrorImportacionExcel("No se pudo leer el archivo — asegúrate de que sea un .xlsx válido.") from exc

    hoja = libro[HOJA_TRABAJADORES] if HOJA_TRABAJADORES in libro.sheetnames else libro.active

    contratistas_por_nombre = {
        _sin_acentos(nombre): id_
        for id_, nombre in EmpresaContratista.objects.filter(activa=True).values_list("id", "nombre")
    }
    encabezados, claves_cursos, claves_certificaciones = _encabezados_y_catalogos()
    indice_cursos = len(ENCABEZADOS_BASE)
    indice_certificaciones = indice_cursos + len(claves_cursos)

    filas = []
    errores_lectura = []
    for numero_fila, fila in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
        if fila is None or all(c in (None, "") for c in fila):
            continue
        celdas = list(fila) + [None] * (len(encabezados) - len(fila))
        celdas = celdas[: len(encabezados)]
        (
            contratista_txt,
            nombres,
            apellidos,
            documento,
            autorizacion_txt,
            eps,
            arl,
            afp,
            vinculacion_txt,
            fecha_inicio,
            fecha_examen,
            fecha_certificacion,
        ) = celdas[:indice_cursos]
        celdas_cursos = celdas[indice_cursos:indice_certificaciones]
        celdas_certificaciones = celdas[indice_certificaciones:]

        if not any([contratista_txt, nombres, apellidos, documento]):
            continue

        contratista_id = contratistas_por_nombre.get(_sin_acentos(contratista_txt or ""))
        if contratista_id is None:
            errores_lectura.append(
                {
                    "fila": numero_fila,
                    "mensaje": (
                        f'No se encontró la empresa contratista "{contratista_txt or "(vacío)"}" — '
                        "usa el desplegable de la columna Contratista."
                    ),
                }
            )
            continue

        autorizacion_datos = _sin_acentos(autorizacion_txt or "") in _VALORES_SI
        vinculacion = _ETIQUETA_A_TIPO_VINCULACION.get(
            _sin_acentos(vinculacion_txt or ""), Trabajador.TipoVinculacion.FIJO
        )

        cursos_safety_academy = {}
        for clave, celda_curso in zip(claves_cursos, celdas_cursos):
            fecha_curso = _valor_fecha(celda_curso)
            if fecha_curso is not None:
                cursos_safety_academy[clave] = fecha_curso.isoformat()

        certificaciones_especiales = {}
        for clave, celda_cert in zip(claves_certificaciones, celdas_certificaciones):
            fecha_cert = _valor_fecha(celda_cert)
            if fecha_cert is not None:
                certificaciones_especiales[clave] = fecha_cert.isoformat()

        filas.append(
            {
                "_fila_excel": numero_fila,
                "contratista": contratista_id,
                "nombres": _texto_celda(nombres),
                "apellidos": _texto_celda(apellidos),
                "documento": _texto_celda(documento),
                "autorizacion_datos": autorizacion_datos,
                "eps": _texto_celda(eps),
                "arl": _texto_celda(arl),
                "afp": _texto_celda(afp),
                "tipo_vinculacion": vinculacion,
                "fecha_inicio_contrato": _valor_fecha(fecha_inicio),
                "fecha_vencimiento_examen_medico": _valor_fecha(fecha_examen),
                "fecha_vencimiento_certificacion_alturas": _valor_fecha(fecha_certificacion),
                "cursos_safety_academy": cursos_safety_academy,
                "certificaciones_especiales": certificaciones_especiales,
            }
        )

    return filas, errores_lectura
